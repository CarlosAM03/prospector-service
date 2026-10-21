"""Reproducible local/CI gate. Dependency preparation is separate from offline checks."""

import argparse
import importlib.util
import json
import os
import subprocess
import sys
import venv
from pathlib import Path

# abspath preserves a Windows subst drive used to work around MAX_PATH.
PACKAGE = Path(os.path.abspath(__file__)).parents[1]


def run(*args: str, cwd: Path = PACKAGE) -> None:
    shown = (*args[: args.index("-c") + 1], "<installed-wheel smoke>") if "-c" in args else args
    print("+", *shown, flush=True)
    subprocess.run(args, cwd=cwd, check=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--wheelhouse",
        type=Path,
        required=True,
        help="Pre-downloaded openpyxl and et_xmlfile wheels (offline install)",
    )
    parser.add_argument(
        "--work-dir",
        type=Path,
        help="Explicit private temporary directory for long Windows source paths",
    )
    options = parser.parse_args()
    for dependency in ("openpyxl", "fastapi", "httpx", "pytest", "ruff", "build"):
        if importlib.util.find_spec(dependency) is None:
            raise RuntimeError(f"complete validation requires installed dependency: {dependency}")
    wheelhouse = options.wheelhouse.absolute()
    if not wheelhouse.is_dir():
        raise ValueError("wheelhouse must exist")
    work = options.work_dir.absolute() if options.work_dir else PACKAGE / ".validation"
    work.mkdir(exist_ok=True, parents=True)
    # Keep subprocess temp writes within this package; don't share Engine caches.
    temporary = work / "tmp"
    temporary.mkdir(exist_ok=True)
    os.environ["TMP"] = os.environ["TEMP"] = os.environ["TMPDIR"] = str(temporary)
    python = sys.executable
    run(python, "-m", "compileall", "-q", "src")
    run(python, "-m", "ruff", "check", ".")
    run(python, "-m", "ruff", "format", "--check", ".")
    run(python, "-m", "pytest", "-q", "-m", "not integration")
    run(python, "-m", "pytest", "-q", "-m", "integration")
    # --no-isolation uses the build requirements installed in the dev environment.
    run(python, "-m", "build", "--no-isolation")
    wheel = PACKAGE / "dist" / "prospector_export-0.1.0-py3-none-any.whl"
    if not wheel.is_file():
        raise RuntimeError("expected wheel was not built")
    # A new path on every run prevents a previous XLSX install tainting CSV-only evidence.
    import tempfile

    with tempfile.TemporaryDirectory(prefix="wheel-", dir=work) as name:
        root = Path(name)
        environment = root / "env"
        venv.EnvBuilder(with_pip=True).create(environment)
        isolated = environment / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        outside = root / "consumer"
        outside.mkdir()
        run(
            str(isolated),
            "-m",
            "pip",
            "install",
            "--no-index",
            "--no-deps",
            str(wheel),
            cwd=outside,
        )
        common = """
import importlib.util, io, csv, pathlib, sys
from prospector_export import ExportService, ExportRecord, MissingExportDependency
import prospector_export
assert pathlib.Path(prospector_export.__file__).is_relative_to(pathlib.Path(sys.prefix))
records = [ExportRecord('CAFÉ 東京', phone=None)]
artifact = ExportService.export(records, 'csv')
assert list(csv.reader(io.StringIO(artifact.content.decode('utf-8'))))[1][0] == 'CAFÉ 東京'
assert 'fastapi' not in sys.modules
"""
        csv_only = (
            common
            + """
assert importlib.util.find_spec('openpyxl') is None
assert importlib.util.find_spec('fastapi') is None
try:
    ExportService.export(records, 'xlsx')
except MissingExportDependency:
    pass
else:
    raise AssertionError('missing XLSX dependency did not fail')
print('CSV_ONLY_INSTALLED_WHEEL_PASS')
"""
        )
        run(str(isolated), "-I", "-c", csv_only, cwd=outside)
        run(
            str(isolated),
            "-m",
            "pip",
            "install",
            "--no-index",
            "--find-links",
            str(wheelhouse),
            "openpyxl>=3.1.5,<4",
            cwd=outside,
        )
        xlsx = (
            common
            + """
from openpyxl import load_workbook
artifact = ExportService.export(records, 'xlsx')
book = load_workbook(artifact.open_binary(), read_only=True)
try:
    assert book.sheetnames == ['Businesses']
    assert list(book['Businesses'].values)[1][0] == 'CAFÉ 東京'
finally:
    book.close()
print('XLSX_INSTALLED_WHEEL_PASS')
"""
        )
        run(str(isolated), "-I", "-c", xlsx, cwd=outside)
    result = {
        "python": sys.version,
        "compile": "PASS",
        "lint": "PASS",
        "format": "PASS",
        "unit": "PASS",
        "integration": "PASS",
        "wheel_build": "PASS",
        "isolated_csv_without_openpyxl": "PASS",
        "isolated_xlsx": "PASS",
    }
    (work / "gate-result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
