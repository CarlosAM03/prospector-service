"""Install prepared wheels offline into a clean venv and exercise the installed application."""

import argparse
import os
import subprocess
import tempfile
import venv
from pathlib import Path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--wheelhouse", type=Path, required=True)
    parser.add_argument("--work-dir", type=Path, default=Path(".validation"))
    options = parser.parse_args()
    wheelhouse = options.wheelhouse.absolute()
    wheels = sorted(str(wheel) for wheel in wheelhouse.glob("*.whl"))
    if not any(Path(wheel).name.startswith("prospector_service-") for wheel in wheels):
        raise SystemExit("Build the Service and dependency wheels before the offline smoke")
    options.work_dir.mkdir(exist_ok=True, parents=True)
    environment = dict(os.environ)
    environment.pop("PYTHONPATH", None)
    environment.pop("PYTHONHOME", None)
    with tempfile.TemporaryDirectory(prefix="installed-", dir=options.work_dir) as name:
        root = Path(name).absolute()
        venv.EnvBuilder(with_pip=True).create(root / "venv")
        python = root / "venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        subprocess.run(
            [str(python), "-m", "pip", "install", "--no-index", "--no-deps", *wheels],
            check=True,
            env=environment,
        )
        consumer = root / "outside"
        consumer.mkdir()
        code = """
from pathlib import Path
import sys
import prospector_service, prospector_engine, prospector_export
from fastapi.testclient import TestClient
from prospector_service.main import create_app
from prospector_service.config import Settings
from prospector_service.adapters.export import export_results
from prospector_service.schemas import BusinessResult
from prospector_engine import ProspectorEngine, EngineConfig, SearchQuery, Source, ProspectorConfigurationError
from openpyxl import load_workbook
for package in (prospector_service, prospector_engine, prospector_export):
    assert Path(package.__file__).is_relative_to(Path(sys.prefix))
assert ProspectorEngine(EngineConfig()).config.limit == 50
try:
    ProspectorEngine(EngineConfig(limit=101)).search(SearchQuery(Source.GOOGLE_MAPS, 'cafe', 'TJ'))
except ProspectorConfigurationError:
    pass
else:
    raise AssertionError('real public Engine preflight did not reject invalid limit')
records = [BusinessResult(name='CAFE UNO', source='google_maps')]
assert b'CAFE UNO' in export_results(records, 'csv').content
book = load_workbook(export_results(records, 'xlsx').open_binary(), read_only=True)
try:
    assert book.sheetnames == ['Businesses']
finally:
    book.close()
with TestClient(create_app(Settings(app_env='test'))) as client:
    assert client.get('/health/live').status_code == 200
    assert client.get('/health/ready').status_code == 503
print('INSTALLED_SERVICE_ENGINE_EXPORT_SMOKE_PASS')
"""
        subprocess.run([str(python), "-I", "-c", code], check=True, cwd=consumer, env=environment)


if __name__ == "__main__":
    main()
