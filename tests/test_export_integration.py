import csv
import io

import pytest
from openpyxl import load_workbook
from prospector_export import HEADERS

from prospector_service.adapters.export import export_results
from prospector_service.adapters.results import map_results

from .test_engine_and_results import query, result


@pytest.mark.integration
@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_export_uses_existing_package_and_normalized_selection(format):
    engine_result = result()
    rows = map_results(engine_result, query())
    artifact = export_results([rows[1], rows[0], rows[1]], format)
    if format == "csv":
        values = list(csv.reader(io.StringIO(artifact.content.decode("utf-8"))))
    else:
        book = load_workbook(artifact.open_binary(), read_only=True)
        try:
            assert book.sheetnames == ["Businesses"]
            values = [list(row) for row in book["Businesses"].values]
        finally:
            book.close()
    assert values[0] == list(HEADERS)
    assert [row[0] for row in values[1:]] == ["CAFE DOS", "CAFE UNO", "CAFE DOS"]
    assert engine_result.original_businesses[0].name == "Cafe Uno"
    assert len(values[0]) == 7


@pytest.mark.integration
def test_empty_export_from_service_adapter():
    assert export_results([], "csv").content.decode().strip() == ",".join(HEADERS)
